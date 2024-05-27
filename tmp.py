from openpyxl import load_workbook

# Load the workbook
file_path = '/home/user/Invoices.xlsx'
wb = load_workbook(file_path)
ws_new = wb['Sheet2']

# Print the contents of "Sheet2"
for row in ws_new.iter_rows(values_only=True):
    print(row)
